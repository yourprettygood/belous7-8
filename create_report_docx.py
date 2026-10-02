import os
import docx
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = docx.Document()

# Page setup: A4, standard margins (left 2.5cm, right 1.5cm, top 2cm, bottom 2cm)
for section in doc.sections:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

# Style setup: Times New Roman, 12pt, 1.5 line spacing, Justified
normal_style = doc.styles['Normal']
normal_style.font.name = 'Times New Roman'
normal_style.font.size = Pt(12)
normal_style.font.color.rgb = RGBColor(0, 0, 0)

# Set rFonts for Russian Cyrillic characters in Word
rFonts = normal_style.element.rPr.get_or_add_rFonts()
rFonts.set(qn('w:ascii'), 'Times New Roman')
rFonts.set(qn('w:hAnsi'), 'Times New Roman')
rFonts.set(qn('w:cs'), 'Times New Roman')
rFonts.set(qn('w:eastAsia'), 'Times New Roman')

def add_para(text="", bold_prefix="", first_line_indent=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    if first_line_indent > 0:
        p.paragraph_format.first_line_indent = Cm(first_line_indent)
    else:
        p.paragraph_format.first_line_indent = Cm(0)
        
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = 'Times New Roman'
        r_bold.font.size = Pt(12)
        r_bold.font.bold = True
        
    if text:
        r_text = p.add_run(text)
        r_text.font.name = 'Times New Roman'
        r_text.font.size = Pt(12)
        r_text.font.bold = False
        
    return p

def add_heading(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.first_line_indent = Cm(1.5)
    
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.bold = True
    return p

def add_code_block(code_text):
    for line in code_text.strip().split('\n'):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Cm(2.0)
        p.paragraph_format.first_line_indent = Cm(0)
        
        r = p.add_run(line)
        r.font.name = 'Courier New'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(30, 30, 30)

def add_image_with_caption(img_path, caption_text):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.paragraph_format.first_line_indent = Cm(0)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.0))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(8)
        p_cap.paragraph_format.first_line_indent = Cm(0)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(11)
        r_cap.font.italic = True

# ================= DOCUMENT CONTENT =================

# Header / Title
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.line_spacing = 1.5
p_title.paragraph_format.space_after = Pt(12)
p_title.paragraph_format.first_line_indent = Cm(0)
r_title = p_title.add_run("Практическое занятие №8-9. Разработка модулей проекта (командная работа). Интеграция модулей проекта (командная работа)")
r_title.font.name = 'Times New Roman'
r_title.font.size = Pt(14)
r_title.font.bold = True

add_para(text=" закрепление практических навыков работы с языком C#, средой программирования Visual Studio, системой управления базами данных PostgreSQL 18 и средой администрирования pgAdmin 4, проведение интеграции программных модулей пользовательского интерфейса, бизнес-логики и базы данных.",
         bold_prefix="Цель занятия:")

add_para(text=" персональный компьютер, среда программирования Visual Studio / .NET 8 SDK, СУБД PostgreSQL 18, инструмент администрирования pgAdmin 4.",
         bold_prefix="Оборудование, технические и программные средства:")

add_para(text=" 2 часа.",
         bold_prefix="Продолжительность занятия:")

add_heading("Задание:")
add_para(text="1. Разработать базу данных, включающую в себя таблицу «Пользователи» и таблицу «Секретный вопрос», используя среду СУБД PostgreSQL и инструмент администрирования pgAdmin 4.")
add_para(text="2. Создать настольное клиентское приложение с окном авторизации и регистрации пользователя на языке C# (платформа WPF).")
add_para(text="3. Создать модель данных на основе разработанной базы данных с использованием технологии ORM Entity Framework Core (библиотека Npgsql.EntityFrameworkCore.PostgreSQL).")
add_para(text="4. Сохранить данные пользователя в созданной базе данных с валидацией введенных данных и проверкой условий пароля.")
add_para(text="5. Создать отчет о проделанной работе в соответствии с установленными требованиями к оформлению.")

add_heading("Теоретические сведения:")
add_para(text="Интеграция программных модулей представляет собой процесс соединения и координации раздельно разработанных программных компонентов (модулей) в единую целостную информационную систему. Основной целью интеграции является обеспечение надежного и согласованного обмена данными между компонентами при сохранении слабой связанности (low coupling) и высокой функциональной сплоченности (high cohesion).")
add_para(text="В современной разработке программных комплексов с базами данных широкое распространение получил архитектурный шаблон N-уровневого приложения, включающий уровень представления (Presentation Layer на базе WPF/XAML), уровень бизнес-логики и доступа к данным (Data Access Layer на базе ORM Entity Framework Core) и уровень постоянного хранения данных (СУБД PostgreSQL). Использование ORM позволяет абстрагироваться от низкоуровневых операций с сокетами и SQL-запросами, оперируя строго типизированными классами языка C#, при этом гарантируя надежную параметризацию запросов и защиту от SQL-инъекций.")

add_heading("Ход работы:")

add_para(text="1. Запуск СУБД PostgreSQL и подключение к серверу через pgAdmin 4. В качестве сервера СУБД используется PostgreSQL 18 (служба postgresql-x64-18), работающая на локальном порту 5432. Для администрирования запущен pgAdmin 4. Было установлено соединение с локальным узлом 127.0.0.1 под учетной записью администратора postgres.")

add_image_with_caption(
    r"d:\GhubReps\belous7-8\docs\images\image1_pgadmin_conn.png",
    "Рис. 1 – Соединение с сервером PostgreSQL в среде pgAdmin 4"
)

add_para(text="2. Создание базы данных и проектирование сущностей. В СУБД была создана новая база данных с именем «user_auth_db». Для реализации хранения сведений об учетных записях пользователей и их параметрах безопасности были спроектированы две взаимосвязанные сущности:")
add_para(text="• Таблица «СекретныйВопрос» (хранит справочник контрольных вопросов, используемых для верификации и восстановления доступа). Атрибуты: «КодСекретногоВопроса» (первичный ключ, целочисленный тип с автоматическим приращением GENERATED ALWAYS AS IDENTITY) и «СекретныйВопрос1» (строка переменной длины VARCHAR(255), содержащая текст вопроса).")
add_para(text="• Таблица «Пользователь» (хранит персональные данные пользователей системы). Атрибуты: «КодПользователя» (первичный ключ, автоинкремент), «Фамилия», «Имя», «ЭлектроннаяПочта», «Пароль», «КодовоеСлово», «ОтветНаСекретныйВопрос», а также внешний ключ «КодСекретногоВопроса», ссылающийся на соответствующую запись таблицы секретных вопросов.")

add_image_with_caption(
    r"d:\GhubReps\belous7-8\docs\images\image2_db_diagram.png",
    "Рис. 2 – Диаграмма структуры базы данных user_auth_db"
)

add_para(text="Для автоматического создания схемы данных в pgAdmin 4 был подготовлен и выполнен следующий DDL-скрипт:")

add_code_block("""-- Создание таблицы СекретныйВопрос
CREATE TABLE "СекретныйВопрос" (
    "КодСекретногоВопроса" INT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    "СекретныйВопрос1" VARCHAR(255) NOT NULL
);

-- Создание таблицы Пользователь с внешним ключом
CREATE TABLE "Пользователь" (
    "КодПользователя" INT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    "Фамилия" VARCHAR(100) NOT NULL,
    "Имя" VARCHAR(100) NOT NULL,
    "ЭлектроннаяПочта" VARCHAR(150) NOT NULL,
    "Пароль" VARCHAR(255) NOT NULL,
    "КодовоеСлово" VARCHAR(100) NOT NULL,
    "ОтветНаСекретныйВопрос" VARCHAR(255) NOT NULL,
    "КодСекретногоВопроса" INT NOT NULL,
    CONSTRAINT "FK_Пользователь_СекретныйВопрос" 
        FOREIGN KEY ("КодСекретногоВопроса") 
        REFERENCES "СекретныйВопрос" ("КодСекретногоВопроса") 
        ON DELETE RESTRICT ON UPDATE CASCADE
);""")

add_para(text="Таблица «СекретныйВопрос» была предварительно заполнена стандартным набором контрольных вопросов (девичья фамилия матери, кличка питомца, любимый город детства, номер первой школы и др.).")

add_para(text="3. Разработка графического интерфейса пользователя (WPF). В среде Visual Studio был создан проект BelousAuthApp на платформе C# (.NET 8 WPF). В файле разметки MainWindow.xaml было спроектировано окно авторизации и регистрации пользователя в строгом соответствии с макетом из методического пособия.")
add_para(text="Графический интерфейс разделен на функциональные зоны: поля персональной информации (Фамилия, Имя, E-mail/логин), поля ввода и повтора пароля с блоком динамической валидации условий надежности (минимальная длина 6 символов, наличие цифр, строчных и заглавных букв, совпадение значений), поле кодового слова, выпадающий список выбора секретного вопроса, текстовое поле ответа на вопрос, флаг согласия с условиями и стилизованная градиентная кнопка «ДАЛЕЕ».")

add_image_with_caption(
    r"d:\GhubReps\belous7-8\docs\images\image3_wpf_auth.png",
    "Рис. 3 – Окно авторизации и регистрации пользователя"
)

add_para(text="4. Создание модели данных и настройка Entity Framework Core. Для осуществления взаимодействия объектной модели приложения на C# и реляционной схемы PostgreSQL был подключен пакет Npgsql.EntityFrameworkCore.PostgreSQL. Были описаны классы сущностей «Пользователь» и «СекретныйВопрос», а также класс контекста «AppDbContext», наследующий от Microsoft.EntityFrameworkCore.DbContext.")

add_image_with_caption(
    r"d:\GhubReps\belous7-8\docs\images\image4_ef_model.png",
    "Рис. 4 – Архитектура модели данных Entity Framework Core"
)

add_para(text="5. Реализация программной логики передачи и сохранения данных. При возникновении события Window_Loaded выполняется LINQ-запрос к контексту базы данных для извлечения актуального списка вопросов: qwest.ItemsSource = context.СекретныйВопрос.Select(i => i.СекретныйВопрос1).ToList().")
add_para(text="По нажатию кнопки «ДАЛЕЕ» вызывается обработчик SignIn_Click, который осуществляет валидацию входных данных, находит идентификатор выбранного секретного вопроса, создает экземпляр класса «Пользователь», помещает его в DbSet через context.Пользователь.Add() и фиксирует транзакцию методом context.SaveChanges(). После успешного добавления выводится информационное окно MessageBox и открывается окно просмотра зарегистрированных пользователей UserWin.")

add_image_with_caption(
    r"d:\GhubReps\belous7-8\docs\images\image5_csharp_code.png",
    "Рис. 5 – Программный код передачи и сохранения данных в C#"
)

add_para(text="6. Реализация окна просмотра зарегистрированных пользователей UserWin. Окно содержит компонент DataGrid, связанный с набором данных context.Пользователь.Include(u => u.СекретныйВопрос).ToList(). Пользовательский интерфейс позволяет просматривать сформированный список пользователей, обновлять данные и переходить к форме регистрации.")

add_image_with_caption(
    r"d:\GhubReps\belous7-8\docs\images\image6_userwin.png",
    "Рис. 6 – Окно просмотра зарегистрированных пользователей (UserWin)"
)

add_para(text="7. Верификация корректности сохранения данных в pgAdmin 4. С помощью встроенного инструмента Query Tool был выполнен контрольный SQL-запрос выборки данных с объединением таблиц Пользователь и СекретныйВопрос. Все внесенные через интерфейс WPF записи были успешно сохранены в реляционную таблицу PostgreSQL с сохранением ссылочной целостности.")

add_image_with_caption(
    r"d:\GhubReps\belous7-8\docs\images\image7_pgadmin_query.png",
    "Рис. 7 – Проверка сохраненных записей в pgAdmin 4 через Query Tool"
)

add_heading("Вывод:")
add_para(text="В ходе выполнения практического занятия №8-9 были успешно закреплены практические навыки проектирования и развертывания реляционных баз данных в СУБД PostgreSQL 18 с использованием графической среды администрирования pgAdmin 4, а также навыки создания настольного клиентского приложения на языке C# (WPF, .NET 8). Была практически реализована интеграция модулей пользовательского интерфейса, бизнес-логики валидации пароля и слоя доступа к данным на основе ORM Entity Framework Core (провайдер Npgsql). Проведенное тестирование подтвердило корректность сохранения учетных записей пользователей и надежность ссылочной целостности базы данных.")

add_heading("Ответы на контрольные вопросы:")

add_para(text="Интеграция программных модулей — это процесс объединения раздельно разработанных программных компонентов (пользовательского интерфейса, бизнес-логики, сервисов доступа к данным, внешних библиотек и СУБД) в единую согласованную программную систему, обеспечивающую их правильное совместное функционирование и информационный обмен по четко определенным интерфейсам и протоколам.",
         bold_prefix="1. Что такое интеграция? ")

add_para(text="Создание базы данных реализуется на уровне СУБД посредством выполнения команд языка определения данных DDL (Data Definition Language), таких как CREATE DATABASE, CREATE TABLE с заданием типов данных столбцов, первичных ключей (PRIMARY KEY), ограничений целостности (NOT NULL, UNIQUE, CHECK) и внешних ключей (FOREIGN KEY). Создание может производиться вручную в интерактивной среде администрирования (pgAdmin 4, SSMS), через консольную утилиту (psql), либо программным путем через миграции ORM (например, методы Database.EnsureCreated() или миграции dotnet ef).",
         bold_prefix="2. Как реализовывается создание базы данных? ")

add_para(text="Модель данных — это логическое и концептуальное описание структуры данных, взаимосвязей между ними, ограничений целостности и правил их обработки в рамках предметной области. В контексте объектно-реляционного отображения (ORM Entity Framework Core) модель данных представляет собой совокупность C#-классов-сущностей (Entity Classes), контекста базы данных (DbContext) и правил их сопоставления (маппинга) с таблицами и связями реляционной СУБД.",
         bold_prefix="3. Что такое модель данных? ")

add_para(text="В классической среде MS SQL Server передача данных из C# осуществляется через клиентский протокол TDS (Tabular Data Stream) с использованием драйвера Microsoft.Data.SqlClient или технологии Entity Framework 6 / EF Core. В архитектуре с СУБД PostgreSQL взаимодействие реализуется по клиент-серверному протоколу PostgreSQL через ADO.NET драйвер Npgsql и провайдер Npgsql.EntityFrameworkCore.PostgreSQL. При вызове метода context.SaveChanges() ORM Entity Framework Core анализирует трекер состояний сущностей (Change Tracker), генерирует параметризованные SQL-инструкции INSERT / UPDATE / DELETE и передает их в бинарном протоколе в СУБД в рамках атомарной транзакции, обеспечивая высокую скорость выполнения, целостность и защиту от SQL-инъекций.",
         bold_prefix="4. Как передаются данные в MSSQL Server и PostgreSQL? ")

output_docx = r"d:\GhubReps\belous7-8\docs\Отчет_ПР_8-9_Белоус.docx"
doc.save(output_docx)
print(f"Report DOCX generated successfully at: {output_docx}")
